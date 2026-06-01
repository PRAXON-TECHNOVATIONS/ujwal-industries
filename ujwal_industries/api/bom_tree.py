import base64
import io
import json

import frappe
from frappe.utils.xlsxutils import make_xlsx

SKIP_FIELD_TYPES = {
	"Section Break", "Column Break", "HTML", "Button", "Image",
	"Fold", "Heading", "Tab Break", "Break", "Table", "Table MultiSelect",
	"Attach", "Attach Image", "Barcode", "Signature",
}

PRIMARY_CHILD = "BOM Item"

TREE_FIELDS = [
	{"fieldname": "level",       "label": "Level"},
	{"fieldname": "parent_item", "label": "Parent Item"},
	{"fieldname": "parent_bom",  "label": "Parent BOM"},
]

DEFAULT_COLUMNS = ["item_code", "item_name", "bom_no", "qty", "uom", "level", "parent_item"]

_SYSTEM_LABELS = {
	"name": "ID", "owner": "Created By", "creation": "Created On",
	"modified_by": "Modified By", "modified": "Last Modified",
	"docstatus": "Document Status", "idx": "Row #", "parent": "Parent BOM Ref",
}

# System/identity fields that must never be updated via import
_NEVER_EDITABLE = frozenset({
	"name", "parent", "parenttype", "parentfield",
	"owner", "creation", "modified_by", "modified", "docstatus", "idx",
})

# module-level cache cleared at start of each export call
_child_cache = {}


def _safe(doctype):
	return doctype.replace(" ", "_")


def _ctbl_prefix(doctype):
	return f"ctbl__{_safe(doctype)}__"


# ── Treeview API ──────────────────────────────────────────────────────────────

@frappe.whitelist()
def get_children(parent=None, is_root=False, selected_boms=None, **filters):
	if not parent or parent == "BOM":
		if selected_boms:
			return _get_specific_boms([b.strip() for b in selected_boms.split(",") if b.strip()])
		return _get_top_level_boms()

	from erpnext.manufacturing.doctype.bom.bom import get_children as erpnext_get_children
	try:
		return erpnext_get_children(parent=parent, is_root=is_root, **filters)
	except StopIteration:
		# StopIteration occurs when a BOM Item's item_code no longer exists in Item.
		# Fall back to our crash-safe implementation for this BOM.
		return _get_bom_children(parent)


@frappe.whitelist()
def get_top_level_boms_list():
	rows = _get_top_level_boms()
	return [{"bom": r.value, "item_code": r.item_code, "item_name": r.item_name} for r in rows]


# ── Export field discovery ────────────────────────────────────────────────────

@frappe.whitelist()
def get_all_export_fields():
	SYSTEM_PARENT = [
		{"fieldname": "name",        "label": "ID",              "fieldtype": "Data"},
		{"fieldname": "owner",       "label": "Created By",      "fieldtype": "Data"},
		{"fieldname": "creation",    "label": "Created On",      "fieldtype": "Datetime"},
		{"fieldname": "modified_by", "label": "Modified By",     "fieldtype": "Data"},
		{"fieldname": "modified",    "label": "Last Modified",   "fieldtype": "Datetime"},
		{"fieldname": "docstatus",   "label": "Document Status", "fieldtype": "Int"},
	]
	SYSTEM_CHILD = SYSTEM_PARENT + [
		{"fieldname": "idx",    "label": "Row #",          "fieldtype": "Int"},
		{"fieldname": "parent", "label": "Parent BOM Ref", "fieldtype": "Data"},
	]

	def _sys(prefix, label_prefix, is_child):
		base = SYSTEM_CHILD if is_child else SYSTEM_PARENT
		return [{"fieldname": prefix + f["fieldname"], "label": label_prefix + f["label"],
		         "fieldtype": f["fieldtype"]} for f in base]

	def _from_meta(doctype, prefix="", label_prefix="", is_child=False):
		out = _sys(prefix, label_prefix, is_child)
		for f in frappe.get_meta(doctype).fields:
			if f.fieldtype in SKIP_FIELD_TYPES or not f.fieldname or not f.label:
				continue
			out.append({"fieldname": prefix + f.fieldname,
			            "label": label_prefix + f.label, "fieldtype": f.fieldtype})
		return out

	bom_meta = frappe.get_meta("BOM")
	groups = [
		{"label": "BOM Item (Items Child Table)",
		 "fields": _from_meta(PRIMARY_CHILD, label_prefix="BOM Item: ", is_child=True)},
		{"label": "BOM / Tree", "fields": list(TREE_FIELDS)},
		{"label": "BOM (Header)",
		 "fields": _from_meta("BOM", prefix="bom__", label_prefix="BOM: ", is_child=False)},
	]

	for f in bom_meta.fields:
		if f.fieldtype != "Table" or not f.options or f.options == PRIMARY_CHILD:
			continue
		child_dt = f.options
		prefix   = _ctbl_prefix(child_dt)
		# Use child_dt (doctype name) as label prefix — must match _build_label_map
		# so DEFAULT_EXPORT_LABELS and Excel headers all use the same prefix.
		fields = _from_meta(child_dt, prefix=prefix, label_prefix=f"{child_dt}: ", is_child=True)
		if fields:
			groups.append({"label": f"{child_dt} (Child Table)", "fields": fields})

	return groups


# ── Export execution ──────────────────────────────────────────────────────────

@frappe.whitelist()
def export_bom_tree(bom=None, selected_boms=None, columns=None):
	global _child_cache
	_child_cache = {}   # clear cache each call

	if isinstance(columns, str):
		columns = json.loads(columns)
	if not columns:
		columns = DEFAULT_COLUMNS

	tree_col_names = {f["fieldname"] for f in TREE_FIELDS}

	# prefix → child doctype (secondary child tables only)
	bom_meta = frappe.get_meta("BOM")
	prefix_to_doctype = {}
	for f in bom_meta.fields:
		if f.fieldtype == "Table" and f.options and f.options != PRIMARY_CHILD:
			prefix_to_doctype[_ctbl_prefix(f.options)] = f.options

	# Categorise columns
	bom_item_cols   = []
	bom_header_cols = []
	child_table_cols = {}   # {doctype: [bare fieldname, ...]}

	for col in columns:
		matched = None
		for prefix, doctype in prefix_to_doctype.items():
			if col.startswith(prefix):
				matched = (doctype, col[len(prefix):])
				break
		if matched:
			doctype, fn = matched
			child_table_cols.setdefault(doctype, [])
			if fn not in child_table_cols[doctype]:
				child_table_cols[doctype].append(fn)
		elif col.startswith("bom__"):
			bom_header_cols.append(col[5:])
		elif col not in tree_col_names:
			bom_item_cols.append(col)

	label_map = _build_label_map(prefix_to_doctype)
	header    = [label_map.get(c, c) for c in columns]
	rows      = [header]

	if selected_boms:
		tops = _get_specific_boms([b.strip() for b in selected_boms.split(",") if b.strip()])
	elif bom and bom != "BOM":
		tops = _get_specific_boms([bom])
	else:
		tops = _get_top_level_boms()

	blank_row = [""] * len(columns)
	for i, top in enumerate(tops):
		if i > 0:
			rows.append(blank_row)
		bom_doc = frappe.get_cached_doc("BOM", top.value)

		# Root row: FG item + BOM-001's own operations (first op inline, extras below)
		rows.extend(_item_with_child_rows(
			item_code=bom_doc.item, item_name=bom_doc.item_name,
			bom_no=bom_doc.name, qty=bom_doc.quantity, uom=bom_doc.uom,
			level=0, parent_item="", parent_bom="",
			sub_bom=bom_doc.name,          # child table rows come from this BOM
			bom_data=bom_doc.as_dict(),
			columns=columns, bom_item_cols=bom_item_cols, bom_header_cols=bom_header_cols,
			tree_col_names=tree_col_names,
			child_table_cols=child_table_cols, prefix_to_doctype=prefix_to_doctype,
			indent="",
		))

		# Recurse into BOM-001's items
		rows.extend(_traverse_bom(
			bom_name=bom_doc.name, level=1,
			parent_item=bom_doc.item, parent_bom=bom_doc.name,
			columns=columns, bom_item_cols=bom_item_cols, bom_header_cols=bom_header_cols,
			tree_col_names=tree_col_names,
			child_table_cols=child_table_cols, prefix_to_doctype=prefix_to_doctype,
		))

	xlsx_file = make_xlsx(rows, "BOM Tree")
	frappe.response["filename"] = "BOM_Tree_Export.xlsx"
	frappe.response["filecontent"] = xlsx_file.getvalue()
	frappe.response["type"] = "binary"


# ── Core row builder ──────────────────────────────────────────────────────────

def _item_with_child_rows(item_code, item_name, bom_no, qty, uom,
                          level, parent_item, parent_bom,
                          sub_bom,           # the BOM whose child records we show
                          bom_data,          # BOM header dict for bom__ columns
                          columns, bom_item_cols, bom_header_cols, tree_col_names,
                          child_table_cols, prefix_to_doctype, indent,
                          item_dict=None):   # full BOM Item record for arbitrary field access
	"""
	Returns 1 + N rows for a single item:
	  Row 0 : item fields  + first record from each selected child table
	  Row 1+ : blank item  + subsequent records (continuation)

	If sub_bom is None, child table columns are all blank (leaf item / no BOM).
	"""
	# Fetch child records for sub_bom (or empty if no sub_bom)
	child_records = {}   # {doctype: [rec, rec, ...]}
	if sub_bom and child_table_cols:
		for doctype, fieldnames in child_table_cols.items():
			cache_key = (sub_bom, doctype)
			if cache_key not in _child_cache:
				all_fields = list(set(fieldnames) | set(_SYSTEM_LABELS.keys()))
				_child_cache[cache_key] = frappe.get_all(
					doctype, filters={"parent": sub_bom},
					fields=all_fields, order_by="idx",
				)
			child_records[doctype] = _child_cache[cache_key]
	else:
		child_records = {dt: [] for dt in child_table_cols}

	# How many rows do we need (at least 1)
	n_rows = max(max((len(v) for v in child_records.values()), default=0), 1)

	# Helper: get child value for column at row index
	def child_val(col, row_idx):
		for prefix, doctype in prefix_to_doctype.items():
			if col.startswith(prefix):
				fn   = col[len(prefix):]
				recs = child_records.get(doctype, [])
				if row_idx < len(recs):
					return recs[row_idx].get(fn, "") or ""
				return ""
		return ""

	rows = []
	for row_idx in range(n_rows):
		row = []
		for col in columns:
			is_child_col = any(col.startswith(p) for p in prefix_to_doctype)

			if is_child_col:
				row.append(child_val(col, row_idx))
			elif row_idx > 0:
				# Continuation row: all item/tree/header columns are blank
				row.append("")
			else:
				# First row: item fields
				if col == "item_code":               row.append(indent + item_code)
				elif col == "item_name":             row.append(item_name or "")
				elif col == "bom_no":                row.append(bom_no or "")
				elif col == "qty":                   row.append(qty or "")
				elif col == "uom":                   row.append(uom or "")
				elif col == "level":                 row.append(level)
				elif col == "parent_item":           row.append(parent_item)
				elif col == "parent_bom":            row.append(parent_bom)
				elif col.startswith("bom__"):        row.append(bom_data.get(col[5:], "") or "")
				else:                                row.append((item_dict.get(col, "") or "") if item_dict else "")
		rows.append(row)

	return rows


def _traverse_bom(bom_name, level, parent_item, parent_bom,
                  columns, bom_item_cols, bom_header_cols, tree_col_names,
                  child_table_cols, prefix_to_doctype):
	rows  = []
	indent = "  " * level
	sys_fns = set(_SYSTEM_LABELS.keys())
	fetch   = list(set(bom_item_cols) | sys_fns | {"item_code", "item_name", "bom_no", "qty", "uom"})

	items = frappe.get_all(PRIMARY_CHILD, filters={"parent": bom_name},
	                       fields=fetch, order_by="idx")

	bom_hdr_cache = {}

	for item in items:
		sub_bom = item.get("bom_no") or None

		# BOM header columns (bom__ prefix) come from the parent BOM
		if bom_header_cols and bom_name not in bom_hdr_cache:
			bom_hdr_cache[bom_name] = frappe.get_cached_doc("BOM", bom_name).as_dict()
		bom_data = bom_hdr_cache.get(bom_name, {})

		# Item row(s) — child table records come from sub_bom (if any)
		rows.extend(_item_with_child_rows(
			item_code=item.get("item_code") or "",
			item_name=item.get("item_name") or "",
			bom_no=sub_bom or "",
			qty=item.get("qty") or "",
			uom=item.get("uom") or "",
			level=level,
			parent_item=parent_item,
			parent_bom=parent_bom,
			sub_bom=sub_bom,
			bom_data=bom_data,
			columns=columns, bom_item_cols=bom_item_cols, bom_header_cols=bom_header_cols,
			tree_col_names=tree_col_names,
			child_table_cols=child_table_cols, prefix_to_doctype=prefix_to_doctype,
			indent=indent,
			item_dict=item,
		))

		# Recurse into sub-BOM items
		if sub_bom:
			rows.extend(_traverse_bom(
				bom_name=sub_bom, level=level + 1,
				parent_item=item.get("item_code") or "",
				parent_bom=sub_bom,
				columns=columns, bom_item_cols=bom_item_cols, bom_header_cols=bom_header_cols,
				tree_col_names=tree_col_names,
				child_table_cols=child_table_cols, prefix_to_doctype=prefix_to_doctype,
			))

	return rows


# ── Label map ─────────────────────────────────────────────────────────────────

def _build_label_map(prefix_to_doctype=None):
	label_map = {f["fieldname"]: f["label"] for f in TREE_FIELDS}

	for fn, lbl in _SYSTEM_LABELS.items():
		label_map[fn] = "BOM Item: " + lbl
	for f in frappe.get_meta(PRIMARY_CHILD).fields:
		if f.fieldname and f.label:
			label_map[f.fieldname] = "BOM Item: " + f.label

	for fn, lbl in _SYSTEM_LABELS.items():
		label_map["bom__" + fn] = "BOM: " + lbl
	for f in frappe.get_meta("BOM").fields:
		if f.fieldname and f.label:
			label_map["bom__" + f.fieldname] = "BOM: " + f.label

	if prefix_to_doctype:
		for prefix, doctype in prefix_to_doctype.items():
			lbl_prefix = doctype + ": "
			for fn, lbl in _SYSTEM_LABELS.items():
				label_map[prefix + fn] = lbl_prefix + lbl
			for f in frappe.get_meta(doctype).fields:
				if f.fieldname and f.label:
					label_map[prefix + f.fieldname] = lbl_prefix + f.label

	return label_map


# ── BOM list helpers ──────────────────────────────────────────────────────────

def _get_bom_children(parent_bom):
	"""
	Fallback for ERPNext's get_children when it raises StopIteration
	(i.e. a BOM Item references an item_code that no longer exists in Item).
	Uses raw SQL for Item lookup to avoid permission-layer issues.
	"""
	bom_doc = frappe.get_cached_doc("BOM", parent_bom)

	bom_items = frappe.db.get_all(
		"BOM Item",
		fields=["item_code", "bom_no as value", "stock_qty", "qty"],
		filters={"parent": parent_bom},
		order_by="idx",
	)

	item_codes = list({(d.item_code or "").strip() for d in bom_items if d.item_code})

	item_map = {}
	if item_codes:
		ph = ", ".join(["%s"] * len(item_codes))
		rows = frappe.db.sql(
			f"""SELECT name, item_name, stock_uom, image, description, is_sub_contracted_item
			    FROM `tabItem` WHERE name IN ({ph})""",
			tuple(item_codes), as_dict=True,
		)
		item_map = {r.name: r for r in rows}

	result = []
	for bom_item in bom_items:
		# Strip leading/trailing spaces from item_code — some BOM Items have dirty data
		# that causes ERPNext's exact-match next() to raise StopIteration.
		item_code = (bom_item.item_code or "").strip()
		bom_item.item_code = item_code
		item_data = item_map.get(item_code, {
			"image": "", "description": "", "name": item_code,
			"stock_uom": "", "item_name": "", "is_sub_contracted_item": 0,
		})
		bom_item.update(item_data)
		bom_item.parent_bom_qty = bom_doc.quantity
		bom_item.expandable = 0 if bom_item.get("value") in ("", None) else 1
		bom_item.image = frappe.db.escape(bom_item.get("image") or "")
		result.append(bom_item)

	return result


def _get_specific_boms(bom_list):
	if not bom_list:
		return []
	placeholders = ", ".join(["%s"] * len(bom_list))
	return frappe.db.sql(
		f"""SELECT b.name AS `value`, b.item AS item_code, b.item_name,
		           b.quantity AS qty, b.uom AS stock_uom,
		           1 AS expandable, '' AS image, '' AS description
		    FROM `tabBOM` b
		    WHERE b.docstatus = 1 AND b.is_active = 1 AND b.name IN ({placeholders})
		    ORDER BY b.item""",
		tuple(bom_list), as_dict=True,
	)


def _get_top_level_boms():
	return frappe.db.sql(
		"""SELECT b.name AS `value`, b.item AS item_code, b.item_name,
		          b.quantity AS qty, b.uom AS stock_uom,
		          1 AS expandable, '' AS image, '' AS description
		   FROM `tabBOM` b
		   WHERE b.docstatus = 1 AND b.is_active = 1
		   AND b.name NOT IN (
		       SELECT DISTINCT bi.bom_no FROM `tabBOM Item` bi
		       INNER JOIN `tabBOM` b2 ON b2.name = bi.parent
		       WHERE b2.docstatus = 1 AND b2.is_active = 1
		       AND bi.bom_no IS NOT NULL AND bi.bom_no != ''
		   )
		   ORDER BY b.item""",
		as_dict=True,
	)


# ── Import helpers ────────────────────────────────────────────────────────────

_SYSTEM_CHILD_TUPLES = [
	("name", "ID"), ("owner", "Created By"), ("creation", "Created On"),
	("modified_by", "Modified By"), ("modified", "Last Modified"),
	("docstatus", "Document Status"), ("idx", "Row #"), ("parent", "Parent BOM Ref"),
]


def _build_reverse_label_map():
	"""Maps Excel header label → (doctype, fieldname)."""
	rev = {}
	# BOM Item fields (label prefix matches _build_label_map: "BOM Item: ")
	for fn, lbl in _SYSTEM_CHILD_TUPLES:
		rev[f"BOM Item: {lbl}"] = ("BOM Item", fn)
	for f in frappe.get_meta("BOM Item").fields:
		if f.fieldname and f.label and f.fieldtype not in SKIP_FIELD_TYPES:
			rev[f"BOM Item: {f.label}"] = ("BOM Item", f.fieldname)
	# Secondary child tables (label prefix: "{child_doctype}: " — matches _build_label_map)
	for f in frappe.get_meta("BOM").fields:
		if f.fieldtype != "Table" or not f.options or f.options == PRIMARY_CHILD:
			continue
		child_dt = f.options
		pfx = child_dt + ": "
		for fn, lbl in _SYSTEM_CHILD_TUPLES:
			rev[pfx + lbl] = (child_dt, fn)
		for cf in frappe.get_meta(child_dt).fields:
			if cf.fieldname and cf.label and cf.fieldtype not in SKIP_FIELD_TYPES:
				rev[pfx + cf.label] = (child_dt, cf.fieldname)
	return rev


def _parse_import_excel(file_b64):
	"""Parse base64-encoded Excel.

	Each record is one of:
	  {"doctype", "name", "fields", "is_new": False}   – update existing row
	  {"doctype", "parent_bom", "fields", "is_new": True} – create new child row
	"""
	import openpyxl

	data = base64.b64decode(file_b64)
	ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).active
	all_rows = list(ws.iter_rows(values_only=True))
	if len(all_rows) < 2:
		return []

	headers = [str(h) if h is not None else "" for h in all_rows[0]]
	rev_map = _build_reverse_label_map()
	col_map = {i: rev_map[h] for i, h in enumerate(headers) if h in rev_map}

	# Carry-forward: the BOM that owns the child-table records on the current row group.
	bom_ctx = None

	# For UPDATE records we accumulate fields across all occurrences of the same ID.
	# This handles shared SFGs that appear multiple times in the export: if the user
	# left the first occurrence untouched and edited the second, the edit still lands.
	# Later non-empty cell values override earlier ones via dict.update().
	update_acc = {}   # (doctype, name) → {fieldname: value}

	# For CREATE records deduplicate by content so a shared SFG doesn't produce
	# duplicate new rows.  Two creates with identical (doctype, parent, fields)
	# are the same logical operation — only the first is kept.
	seen_create_keys = set()
	create_records = []

	for raw_row in all_rows[1:]:
		cells = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))

		dt_ids = {}
		dt_fields = {}

		for idx, (dt, fn) in col_map.items():
			val = cells[idx]
			if val is None or not str(val).strip():
				continue
			val = val if not isinstance(val, str) else val.strip()
			if fn == "name":
				dt_ids[dt] = str(val)
			else:
				dt_fields.setdefault(dt, {})[fn] = val

		# Blank rows are section separators between top-level BOMs in the export.
		# Reset carry-forward so a new section never inherits the previous BOM context.
		if not dt_ids and not dt_fields:
			bom_ctx = None
			continue

		# Carry-forward step 1: bom_no column gives the sub-BOM hint
		item_bom_no = dt_fields.get("BOM Item", {}).get("bom_no")
		if item_bom_no:
			bom_ctx = str(item_bom_no)

		# Carry-forward step 2: DB parent of any existing child record (most accurate)
		for dt, rec_id in dt_ids.items():
			if dt == "BOM Item":
				continue
			try:
				db_parent = frappe.db.get_value(dt, rec_id, "parent")
				if db_parent:
					bom_ctx = db_parent
					break
			except Exception:
				pass

		# Accumulate field edits for existing records
		for dt, rec_id in dt_ids.items():
			uid = (dt, rec_id)
			fields = {fn: v for fn, v in dt_fields.get(dt, {}).items()
			          if fn not in _NEVER_EDITABLE}
			if fields:
				if uid not in update_acc:
					update_acc[uid] = {}
				update_acc[uid].update(fields)  # later occurrences win

		# Collect new records (no ID)
		for dt, all_fields in dt_fields.items():
			if dt in dt_ids:
				continue
			parent_bom = all_fields.get("parent")
			if not parent_bom and dt != "BOM Item":
				parent_bom = bom_ctx
			if not parent_bom:
				continue
			editable = {fn: v for fn, v in all_fields.items() if fn not in _NEVER_EDITABLE}
			if not editable:
				continue
			# Normalise values to lowercase strings so "2nd Bending" and "2nd bending"
			# (typed inconsistently across duplicate BOM sections) are treated as one.
			create_key = (dt, str(parent_bom),
			              frozenset((fn, str(v).strip().lower()) for fn, v in editable.items()))
			if create_key in seen_create_keys:
				continue
			seen_create_keys.add(create_key)
			create_records.append({"doctype": dt, "parent_bom": str(parent_bom),
			                       "fields": editable, "is_new": True})

	records = [
		{"doctype": dt, "name": rec_id, "fields": fields, "is_new": False}
		for (dt, rec_id), fields in update_acc.items()
		if fields
	]
	records.extend(create_records)
	return records


def _parentfield_for(child_doctype):
	"""Return the fieldname in BOM that holds this child table."""
	for f in frappe.get_meta("BOM").fields:
		if f.fieldtype == "Table" and f.options == child_doctype:
			return f.fieldname
	return None


def _vals_equal(old, new):
	"""Compare two values loosely (handles numeric types from Excel)."""
	try:
		return float(old) == float(new)
	except (ValueError, TypeError):
		return str(old or "").strip() == str(new or "").strip()


# ── Import APIs ───────────────────────────────────────────────────────────────

@frappe.whitelist()
def preview_bom_import(file_b64):
	"""Returns list of change objects (field updates + new record creates)."""
	records = _parse_import_excel(file_b64)
	if not records:
		return {"changes": [], "warning": "no_records"}

	changes = []
	for rec in records:
		if rec.get("is_new"):
			changes.append({
				"change_type": "create",
				"doctype": rec["doctype"],
				"name": None,
				"bom": rec.get("parent_bom", ""),
				"new_fields": rec["fields"],
			})
		else:
			editable_fns = [fn for fn in rec["fields"] if fn not in _NEVER_EDITABLE]
			if not editable_fns:
				continue
			try:
				fetch = list(set(editable_fns) | {"parent"})
				current = frappe.db.get_value(rec["doctype"], rec["name"], fetch, as_dict=True)
			except Exception:
				continue
			if not current:
				continue
			bom = current.get("parent") or ""
			for fn in editable_fns:
				old_val = current.get(fn, "")
				new_val = rec["fields"][fn]
				if not _vals_equal(old_val, new_val):
					changes.append({
						"change_type": "update",
						"doctype": rec["doctype"],
						"name": rec["name"],
						"field": fn,
						"old": old_val,
						"new": new_val,
						"bom": bom,
					})

	# Enrich with BOM item/name for the grouped preview
	bom_names = list({c["bom"] for c in changes if c.get("bom")})
	if bom_names:
		bom_rows = frappe.db.get_all("BOM", filters={"name": ["in", bom_names]},
		                             fields=["name", "item", "item_name"])
		bom_map = {r.name: r for r in bom_rows}
		for c in changes:
			info = bom_map.get(c.get("bom") or "")
			c["bom_item"] = info.item if info else ""
			c["bom_item_name"] = info.item_name if info else ""

	return {"changes": changes, "warning": None}


@frappe.whitelist()
def apply_bom_import(file_b64):
	"""Apply field updates and new record creates from the Excel file."""
	records = _parse_import_excel(file_b64)
	updated = created = skipped = 0
	errors = []

	for rec in records:
		if rec.get("is_new"):
			dt = rec["doctype"]
			parent_bom = rec.get("parent_bom")
			if not parent_bom:
				skipped += 1
				continue
			parentfield = _parentfield_for(dt)
			if not parentfield:
				errors.append(f"Cannot find parentfield for {dt} in BOM")
				skipped += 1
				continue
			try:
				result = frappe.db.sql(
					f"SELECT COALESCE(MAX(idx), 0) FROM `tab{dt}` WHERE parent = %s",
					parent_bom,
				)
				next_idx = (result[0][0] if result else 0) + 1
				doc = frappe.new_doc(dt)
				doc.parent = parent_bom
				doc.parenttype = "BOM"
				doc.parentfield = parentfield
				doc.idx = next_idx
				for fn, val in rec["fields"].items():
					doc.set(fn, val)
				doc.db_insert()
				created += 1
			except Exception as e:
				errors.append(f"New {dt} in {parent_bom}: {e}")
		else:
			to_update = {fn: val for fn, val in rec["fields"].items()
			             if fn not in _NEVER_EDITABLE}
			if not to_update:
				skipped += 1
				continue
			try:
				for fn, val in to_update.items():
					frappe.db.set_value(rec["doctype"], rec["name"], fn, val)
				updated += 1
			except Exception as e:
				errors.append(f"{rec['doctype']} {rec['name']}: {e}")

	frappe.db.commit()
	return {"updated": updated, "created": created, "skipped": skipped, "errors": errors}
